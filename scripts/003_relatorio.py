import argparse
import base64
import json
import re
import subprocess
from datetime import datetime
from pathlib import Path
from html import escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats
import markdown
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


# ============================================================
# CONFIGURAÇÕES
# ============================================================

OPEN_CODE = (
    Path.home()
    / "AppData/Roaming/npm/node_modules/opencode-ai/bin/opencode.exe"
)
MODELO_PADRAO = "opencode/deepseek-v4.1-flash-free"

PASTA_ATUAL = Path(__file__).resolve().parent
PASTA_RAIZ = (
    PASTA_ATUAL.parent
    if PASTA_ATUAL.name in ("runs", "scripts")
    else PASTA_ATUAL
)

ARQUIVO_DADOS_CSV = (
    PASTA_RAIZ / "output" / "01_coleta" / "dados_clima.csv"
)
ARQUIVO_ESTATISTICAS = (
    PASTA_RAIZ / "output" / "02_analise" / "analise_estatistica.json"
)
ARQUIVO_PROMPT = PASTA_RAIZ / "prompts" / "03_prompt_relatorio.md"

PASTA_OUTPUT = PASTA_RAIZ / "output" / "03_relatorio"
PASTA_GRAFICOS = PASTA_OUTPUT / "graficos"
PASTA_WORD = PASTA_OUTPUT / "word"
PASTA_HTML = PASTA_OUTPUT / "html"
PASTA_MD = PASTA_OUTPUT / "markdown"
PASTA_PLANOS = PASTA_OUTPUT / "planos"
PASTA_RUNS = PASTA_RAIZ / "runs"

ARQUIVO_SAIDA_PLANO = PASTA_PLANOS / "plano_editorial.json"
ARQUIVO_SAIDA_MD = PASTA_MD / "relatorio_climatico.md"
ARQUIVO_SAIDA_HTML = PASTA_HTML / "relatorio_climatico.html"
ARQUIVO_SAIDA_DOCX = PASTA_WORD / "relatorio_climatico.docx"

for pasta in (
    PASTA_OUTPUT,
    PASTA_GRAFICOS,
    PASTA_WORD,
    PASTA_HTML,
    PASTA_MD,
    PASTA_PLANOS,
    PASTA_RUNS,
):
    pasta.mkdir(parents=True, exist_ok=True)


# ============================================================
# CONSULTA AO OPENCODE
# ============================================================

def consultar_modelo(texto_prompt, modelo=MODELO_PADRAO):
    """Consulta o OpenCode e apresenta detalhes caso não haja resposta."""

    if not OPEN_CODE.exists():
        raise FileNotFoundError(
            f"Executável OpenCode não encontrado em: {OPEN_CODE}"
        )

    print(f"   Modelo: {modelo}")
    print(f"   Tamanho do prompt: {len(texto_prompt):,} caracteres")
    print("   Aguardando resposta do modelo...")

    try:

        resposta = subprocess.run(
            [
                str(OPEN_CODE),
                "run",
                "--agent", "general",
                "--model", modelo,
                "--dir", str(PASTA_RAIZ),
                texto_prompt,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
        )

    except subprocess.TimeoutExpired:
        raise RuntimeError(
            "O OpenCode excedeu o limite de 300 segundos. "
            "O prompt pode ser grande ou o modelo pode estar lento."
        )

    saida = resposta.stdout.strip()
    erro = resposta.stderr.strip()

    print("\n--- DEBUG OPENCODE ---")
    print("Código:", resposta.returncode)
    print("Tamanho stdout:", len(resposta.stdout))
    print("Tamanho stderr:", len(resposta.stderr))
    print("STDOUT:", repr(resposta.stdout[:2000]))
    print("STDERR:", repr(resposta.stderr[:4000]))
    print("--- FIM DEBUG ---\n")

    print(f"   Código de saída: {resposta.returncode}")
    print(f"   Caracteres em stdout: {len(saida)}")

    if resposta.returncode != 0:
        raise RuntimeError(
            "Falha na execução do OpenCode.\n"
            f"Detalhes: {erro or 'Nenhuma mensagem de erro retornada.'}"
        )

    if not saida:
        raise RuntimeError(
            "O OpenCode terminou sem retornar texto.\n"
            f"Modelo: {modelo}\n"
            f"Prompt: {len(texto_prompt):,} caracteres\n"
            f"stderr: {erro[:5000] or '(vazio)'}"
        )

    return saida


# ============================================================
# LIMPEZA DE TEXTO
# ============================================================

def sanitizar_texto_latex(texto):
    """Converte notações matemáticas frequentes para texto legível."""

    if not texto:
        return ""

    t = texto

    substituicoes = [
        (r"\\,?\^\\circ\\text\{C\}", "°C"),
        (r"\\circ\\text\{C\}", "°C"),
        (r"\\circ C", "°C"),
        (r"\\bar\{x\}", "x̄"),
        (r"\\tilde\{x\}", "Mediana"),
        (r"\\mu", "μ"),
        (r"\\sigma\^2", "s²"),
        (r"\\sigma", "s"),
        (r"\\rho_s", "rho"),
        (r"\\rho", "r"),
        (r"\\beta_1", "β"),
        (r"\\beta_0", "α"),
        (r"\\beta", "β"),
        (r"\\alpha", "α"),
        (r"\\times", "×"),
        (r"\\le", "≤"),
        (r"\\ge", "≥"),
        (r"\\ne", "≠"),
        (r"\\approx", "≈"),
        (r"\\pm", "±"),
        (r"\\sim", "~"),
        (r"\\in", "∈"),
        (r"\\sum", "Soma"),
        (r"\\mathcal\{N\}", "Normal"),
        (r"\\text\{([^}]*)\}", r"\1"),
        (r"\\frac\{([^}]*)\}\{([^}]*)\}", r"(\1) / (\2)"),
    ]

    for padrao, substituicao in substituicoes:
        t = re.sub(padrao, substituicao, t)

    t = re.sub(r"\$\$([^\$]+)\$\$", r"\1", t)
    t = re.sub(r"\$([^\$]+)\$", r"\1", t)
    t = t.replace("\\", "")

    return t


# ============================================================
# EXTRAÇÃO DO PLANO EDITORIAL E DO MARKDOWN
# ============================================================

def extrair_resposta_dupla(resposta_modelo):
    """Extrai o plano JSON e o relatório Markdown da resposta da IA."""

    plano_editorial = None
    texto_markdown = ""

    # Primeiro, tenta extrair um bloco JSON cercado por crases.
    blocos_json = re.findall(
        r"```(?:json)?\s*([\s\S]*?)```",
        resposta_modelo,
        flags=re.IGNORECASE,
    )

    for bloco in blocos_json:
        try:
            candidato = json.loads(bloco.strip())
            if isinstance(candidato, dict) and (
                "figuras_encomendadas" in candidato
            ):
                plano_editorial = candidato
                break
        except json.JSONDecodeError:
            continue

    # Se necessário, procura objetos JSON no texto.
    if plano_editorial is None:
        decoder = json.JSONDecoder()

        for match in re.finditer(r"\{", resposta_modelo):
            try:
                candidato, _ = decoder.raw_decode(
                    resposta_modelo[match.start():]
                )

                if (
                    isinstance(candidato, dict)
                    and "figuras_encomendadas" in candidato
                ):
                    plano_editorial = candidato
                    break

            except json.JSONDecodeError:
                continue

    # Extrai o Markdown de um bloco específico, se houver.
    match_md = re.search(
        r"```(?:markdown|md)\s*([\s\S]*?)```",
        resposta_modelo,
        flags=re.IGNORECASE,
    )

    if match_md:
        texto_markdown = match_md.group(1).strip()
    else:
        # Remove os blocos JSON e preserva o texto restante.
        texto_markdown = resposta_modelo

        texto_markdown = re.sub(
            r"```(?:json)?\s*[\s\S]*?```",
            "",
            texto_markdown,
            flags=re.IGNORECASE,
        ).strip()

    if plano_editorial is None:
        raise ValueError(
            "A resposta da IA não contém um plano editorial JSON válido "
            "com o campo 'figuras_encomendadas'."
        )

    if not texto_markdown:
        raise ValueError(
            "A resposta da IA não contém o texto do relatório Markdown."
        )

    if not isinstance(plano_editorial.get("figuras_encomendadas"), list):
        plano_editorial["figuras_encomendadas"] = []

    return plano_editorial, texto_markdown


# ============================================================
# CONFIGURAÇÃO VISUAL DOS GRÁFICOS
# ============================================================

def configurar_estilo():
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["font.sans-serif"] = [
        "DejaVu Sans",
        "Arial",
        "Helvetica",
    ]
    plt.rcParams["axes.edgecolor"] = "#cbd5e1"
    plt.rcParams["axes.linewidth"] = 0.8
    plt.rcParams["grid.color"] = "#e2e8f0"
    plt.rcParams["grid.linestyle"] = "--"
    plt.rcParams["grid.alpha"] = 0.7


def render_distribuicao(df, params, titulo, caminho_saida):
    """Gera histogramas e boxplots para uma ou duas variáveis."""

    configurar_estilo()

    var1 = params.get("variavel_1") or params.get("variavel")
    var2 = params.get("variavel_2")

    cols = [
        col
        for col in (var1, var2)
        if col and col in df.columns
        and pd.api.types.is_numeric_dtype(df[col])
    ]

    if not cols:
        return None

    n_cols = len(cols)

    fig, axes = plt.subplots(
        2,
        n_cols,
        figsize=(6 * n_cols, 7),
        dpi=200,
        gridspec_kw={"height_ratios": [2.5, 1]},
        squeeze=False,
    )

    cores = ["#3b82f6", "#0ea5e9"]

    for idx, col in enumerate(cols):
        serie = pd.to_numeric(df[col], errors="coerce").dropna()

        if serie.empty:
            continue

        media = serie.mean()
        mediana = serie.median()

        ax_hist = axes[0, idx]
        ax_hist.hist(
            serie,
            bins=min(10, max(1, int(np.sqrt(len(serie))))),
            color=cores[idx % len(cores)],
            alpha=0.65,
            edgecolor="#1e3a8a",
            density=True,
        )

        ax_hist.axvline(
            media,
            color="#dc2626",
            linewidth=2,
            label=f"Média: {media:.2f}",
        )
        ax_hist.axvline(
            mediana,
            color="#16a34a",
            linewidth=2,
            linestyle="--",
            label=f"Mediana: {mediana:.2f}",
        )

        ax_hist.set_title(f"Distribuição: {col}", fontsize=11, fontweight="bold")
        ax_hist.set_ylabel("Densidade")
        ax_hist.legend(frameon=True, facecolor="#f8fafc", fontsize=9)
        ax_hist.grid(True)

        ax_box = axes[1, idx]
        ax_box.boxplot(
            serie,
            vert=False,
            patch_artist=True,
            boxprops=dict(facecolor="#93c5fd", color="#1e3a8a"),
            medianprops=dict(color="#16a34a", linewidth=2),
            widths=0.55,
        )
        ax_box.set_xlabel(col)
        ax_box.set_yticks([])
        ax_box.grid(True)

    fig.suptitle(
        titulo or "Análise Comparativa de Distribuição",
        fontsize=12,
        fontweight="bold",
        y=0.98,
    )
    fig.tight_layout()
    fig.savefig(caminho_saida, bbox_inches="tight")
    plt.close(fig)

    return caminho_saida


def render_dispersao(df, params, titulo, caminho_saida):
    """Gera dispersão e reta OLS quando os dados permitem o cálculo."""

    configurar_estilo()

    vx = params.get("var_x")
    vy = params.get("var_y")

    if not vx or not vy or vx not in df.columns or vy not in df.columns:
        return None

    sub = df[[vx, vy]].copy()

    sub[vx] = pd.to_numeric(sub[vx], errors="coerce")
    sub[vy] = pd.to_numeric(sub[vy], errors="coerce")
    sub = sub.replace([np.inf, -np.inf], np.nan).dropna()

    if len(sub) < 3:
        print(f"      Aviso: dispersão ignorada ({vx}, {vy}): menos de 3 observações.")
        return None

    x = sub[vx].to_numpy()
    y = sub[vy].to_numpy()

    if np.unique(x).size < 2 or np.unique(y).size < 2:
        print(f"      Aviso: dispersão ignorada ({vx}, {vy}): variável constante.")
        return None

    r, p = stats.pearsonr(x, y)
    slope, intercept = np.polyfit(x, y, 1)

    fig, ax = plt.subplots(figsize=(7.5, 5.5), dpi=200)

    ax.scatter(
        x,
        y,
        color="#ea580c",
        s=35 if len(x) > 100 else 70,
        edgecolor="#7c2d12",
        alpha=0.65,
        zorder=3,
    )

    x_grid = np.linspace(x.min(), x.max(), 100)
    ax.plot(
        x_grid,
        slope * x_grid + intercept,
        color="#c2410c",
        linewidth=2,
        linestyle="--",
        label=f"Reta OLS (r = {r:.4f}, p = {p:.4g})",
    )

    ax.set_title(
        titulo or f"Diagrama de Dispersão: {vx} vs. {vy}",
        fontsize=11,
        fontweight="bold",
    )
    ax.set_xlabel(vx)
    ax.set_ylabel(vy)
    ax.legend(frameon=True, facecolor="#f8fafc")
    ax.grid(True)

    fig.tight_layout()
    fig.savefig(caminho_saida, bbox_inches="tight")
    plt.close(fig)

    return caminho_saida


def render_matriz_correlacao(df, params, titulo, caminho_saida):
    """Gera heatmap de correlação de Pearson."""

    configurar_estilo()

    lista_vars = params.get("lista_variaveis", [])

    cols = [
        col for col in lista_vars
        if col in df.columns
        and pd.api.types.is_numeric_dtype(df[col])
    ]

    if len(cols) < 2:
        cols = df.select_dtypes(include=[np.number]).columns[:6].tolist()

    if len(cols) < 2:
        return None

    sub = df[cols].replace([np.inf, -np.inf], np.nan)
    matriz_df = sub.corr(method="pearson", min_periods=2)
    matriz = matriz_df.to_numpy()

    fig, ax = plt.subplots(figsize=(8, 6.5), dpi=200)

    cax = ax.imshow(
        np.ma.masked_invalid(matriz),
        cmap="coolwarm",
        vmin=-1,
        vmax=1,
    )
    fig.colorbar(cax, ax=ax, fraction=0.046, pad=0.04)

    ax.set_xticks(range(len(cols)))
    ax.set_yticks(range(len(cols)))
    ax.set_xticklabels(cols, rotation=35, ha="right", fontsize=9)
    ax.set_yticklabels(cols, fontsize=9)

    for i in range(len(cols)):
        for j in range(len(cols)):
            val = matriz[i, j]

            if np.isfinite(val):
                cor = "white" if abs(val) > 0.55 else "black"
                ax.text(
                    j,
                    i,
                    f"{val:.2f}",
                    ha="center",
                    va="center",
                    color=cor,
                    fontsize=9,
                    fontweight="bold",
                )
            else:
                ax.text(
                    j,
                    i,
                    "N/A",
                    ha="center",
                    va="center",
                    color="black",
                    fontsize=8,
                )

    ax.set_title(
        titulo or "Matriz de Correlação Linear de Pearson",
        fontsize=11,
        fontweight="bold",
        pad=12,
    )

    fig.tight_layout()
    fig.savefig(caminho_saida, bbox_inches="tight")
    plt.close(fig)

    return caminho_saida


def render_serie_temporal(df, params, titulo, caminho_saida):
    """Gera série temporal ordenada cronologicamente e reta de tendência."""

    configurar_estilo()

    var = params.get("variavel") or params.get("var_y")

    if not var or var not in df.columns:
        return None

    colunas = ["data", var] if "data" in df.columns else [var]
    sub = df[colunas].copy()

    sub[var] = pd.to_numeric(sub[var], errors="coerce")
    sub = sub.replace([np.inf, -np.inf], np.nan).dropna(subset=[var])

    if "data" in sub.columns:
        sub["data"] = pd.to_datetime(sub["data"], errors="coerce")
        sub = sub.dropna(subset=["data"]).sort_values("data")

    sub = sub.reset_index(drop=True)

    if len(sub) < 3:
        print(f"      Aviso: série temporal ignorada ({var}): menos de 3 observações.")
        return None

    y = sub[var].to_numpy()

    if np.unique(y).size < 2:
        print(f"      Aviso: série temporal ignorada ({var}): variável constante.")
        return None

    if "data" in sub.columns:
        dias = (
            (sub["data"] - sub["data"].min()).dt.total_seconds()
            / 86400
        ).to_numpy()
    else:
        dias = np.arange(len(sub), dtype=float)

    if np.unique(dias).size < 2:
        return None

    reg = stats.linregress(dias, y)

    fig, ax = plt.subplots(figsize=(8, 4.8), dpi=200)

    ax.plot(
        dias,
        y,
        marker="o",
        color="#2563eb",
        linewidth=1.8,
        label=f"Série {var}",
    )

    ax.plot(
        dias,
        reg.slope * dias + reg.intercept,
        color="#dc2626",
        linestyle="--",
        linewidth=2,
        label=(
            f"Tendência (β = {reg.slope:.4f}/dia, "
            f"R² = {reg.rvalue**2:.4f}, p = {reg.pvalue:.4g})"
        ),
    )

    if "data" in sub.columns:
        datas = sub["data"].dt.strftime("%d/%m")
        ax.set_xticks(dias[::max(1, len(dias) // 10)])
        ax.set_xticklabels(
            datas.iloc[::max(1, len(datas) // 10)],
            rotation=30,
        )
    else:
        ax.set_xlabel("Observação")

    ax.set_title(
        titulo or f"Série Temporal de {var}",
        fontsize=11,
        fontweight="bold",
    )
    ax.set_ylabel(var)
    ax.legend(frameon=True, facecolor="#f8fafc")
    ax.grid(True)

    fig.tight_layout()
    fig.savefig(caminho_saida, bbox_inches="tight")
    plt.close(fig)

    return caminho_saida


def despachar_geracao_figura(df, config_figura, indice, pasta_graficos):
    """Seleciona e executa a ferramenta gráfica solicitada pelo agente."""

    tipo = config_figura.get("tipo_ferramenta", "").lower()
    params = config_figura.get("parametros", {})
    titulo = config_figura.get("titulo_grafico", "")

    nome_arquivo = f"figura_{indice}_{tipo}.png"
    caminho = pasta_graficos / nome_arquivo

    ferramentas = {
        "distribuicao": render_distribuicao,
        "dispersao": render_dispersao,
        "matriz_correlacao": render_matriz_correlacao,
        "serie_temporal": render_serie_temporal,
    }

    ferramenta = ferramentas.get(tipo)

    if ferramenta is None:
        print(f"      Aviso: ferramenta gráfica desconhecida: {tipo}")
        return None

    try:
        return ferramenta(df, params, titulo, caminho)
    except (ValueError, TypeError, KeyError, np.linalg.LinAlgError) as erro:
        print(f"      Aviso: não foi possível gerar {nome_arquivo}: {erro}")
        return None


# ============================================================
# EXPORTAÇÃO PARA HTML
# ============================================================

def converter_imagem_base64(caminho_imagem):
    bytes_img = Path(caminho_imagem).read_bytes()
    b64_str = base64.b64encode(bytes_img).decode("utf-8")
    return f"data:image/png;base64,{b64_str}"


def gerar_relatorio_html(
    texto_markdown,
    figuras_geradas,
    caminho_html_saida,
    titulo_relatorio,
):
    """Gera HTML com o relatório e as figuras incorporadas."""

    texto_limpo = sanitizar_texto_latex(texto_markdown)

    conteudo_html_corpo = markdown.markdown(
        texto_limpo,
        extensions=["tables", "fenced_code"],
    )

    cards_figuras = []

    for fig in figuras_geradas:
        caminho = Path(fig["caminho"])

        if caminho.exists():
            legenda = sanitizar_texto_latex(fig.get("legenda", ""))

            cards_figuras.append(
                f"""
                <div class="card-figure">
                    <img src="{converter_imagem_base64(caminho)}"
                         alt="{escape(legenda, quote=True)}" />
                    <p class="figure-caption">
                        <strong>{escape(fig['id'].capitalize())}:</strong>
                        {escape(legenda)}
                    </p>
                </div>
                """
            )

    galeria_html = "".join(cards_figuras)
    titulo_html = escape(str(titulo_relatorio))

    template_completo = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{titulo_html}</title>
<style>
:root {{
    --primary: #0f172a;
    --secondary: #1e3a8a;
    --accent: #2563eb;
    --text: #1e293b;
    --bg: #f8fafc;
    --border: #e2e8f0;
}}
body {{
    font-family: Arial, Helvetica, sans-serif;
    background: var(--bg);
    color: var(--text);
    line-height: 1.65;
    margin: 0;
    padding: 30px 15px;
}}
.container {{
    max-width: 920px;
    margin: 0 auto;
    background: #fff;
    padding: 45px 55px;
    border-radius: 12px;
    border: 1px solid var(--border);
}}
h1 {{
    color: var(--primary);
    font-size: 26px;
    border-bottom: 3px solid var(--accent);
    padding-bottom: 12px;
}}
h2 {{
    color: var(--secondary);
    font-size: 20px;
    margin-top: 32px;
    border-left: 4px solid var(--accent);
    padding-left: 10px;
}}
h3 {{ color: var(--primary); font-size: 16px; }}
p, li {{ font-size: 15px; color: #334155; }}
table {{
    width: 100%;
    border-collapse: collapse;
    margin: 20px 0;
    font-size: 14px;
}}
th, td {{
    padding: 10px 14px;
    text-align: left;
    border-bottom: 1px solid var(--border);
}}
th {{ background: #f1f5f9; color: var(--primary); }}
blockquote {{
    background: #eff6ff;
    border-left: 4px solid var(--accent);
    margin: 15px 0;
    padding: 12px 18px;
}}
.gallery-section {{
    margin-top: 45px;
    padding-top: 25px;
    border-top: 2px dashed var(--border);
}}
.gallery-title {{
    text-align: center;
    border: none;
    padding: 0;
}}
.card-figure {{
    background: #fff;
    border: 1px solid var(--border);
    border-radius: 8px;
    padding: 15px;
    margin-bottom: 30px;
    text-align: center;
    page-break-inside: avoid;
}}
.card-figure img {{ max-width: 100%; height: auto; }}
.figure-caption {{ font-size: 13px; color: #475569; }}
@media print {{
    body {{ background: #fff; padding: 0; }}
    .container {{ border: none; padding: 15px; max-width: 100%; }}
    .card-figure {{ page-break-inside: avoid; }}
}}
</style>
</head>
<body>
<div class="container">
{conteudo_html_corpo}
<section class="gallery-section">
<h2 class="gallery-title">Painel de Evidências Gráficas</h2>
{galeria_html}
</section>
</div>
</body>
</html>
"""

    caminho_html_saida.write_text(
        template_completo,
        encoding="utf-8",
    )


# ============================================================
# EXPORTAÇÃO PARA WORD
# ============================================================

def aplicar_borda_tabela(tabela):
    """Aplica sombreamento ao cabeçalho da tabela Word."""

    shading_xml = parse_xml(
        r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls("w"))
    )

    for cell in tabela.rows[0].cells:
        cell._tc.get_or_add_tcPr().append(shading_xml)


def gerar_relatorio_docx(
    texto_markdown,
    figuras_geradas,
    caminho_docx_saida,
    metadados,
    titulo_relatorio,
):
    """Gera o relatório Word com formatação acadêmica."""

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(3.0)
        section.left_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(12)
    estilo_normal.font.color.rgb = RGBColor(0, 0, 0)
    estilo_normal.paragraph_format.line_spacing = 1.5
    estilo_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.space_after = Pt(12)

    run_tit = p_titulo.add_run(
        f"{sanitizar_texto_latex(titulo_relatorio).upper()}\n"
    )
    run_tit.bold = True
    run_tit.font.size = Pt(14)

    periodo = metadados.get("periodo", {})

    run_sub = p_titulo.add_run(
        f"Estudo de Caso: {metadados.get('cidade', 'Não informada')} | "
        f"Período: {periodo.get('inicio', '')} a {periodo.get('fim', '')}\n"
    )
    run_sub.font.size = Pt(11)
    run_sub.italic = True

    linhas = texto_markdown.splitlines()
    i = 0

    while i < len(linhas):
        linha_bruta = linhas[i].strip()
        linha = sanitizar_texto_latex(linha_bruta)

        if not linha:
            i += 1
            continue

        if linha.startswith("# "):
            i += 1
            continue

        elif linha.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)

            run = h.add_run(linha[3:].upper())
            run.bold = True
            run.font.size = Pt(12)

        elif linha.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)

            run = h.add_run(linha[4:])
            run.bold = True
            run.font.size = Pt(12)

        elif linha.startswith("|") and "|" in linha[1:]:
            linhas_tabela = []

            while i < len(linhas) and linhas[i].strip().startswith("|"):
                linhas_tabela.append(
                    sanitizar_texto_latex(linhas[i].strip())
                )
                i += 1

            if len(linhas_tabela) >= 2:
                cabecalho = [
                    col.strip()
                    for col in linhas_tabela[0].strip("|").split("|")
                ]

                linhas_dados = []

                for linha_tabela in linhas_tabela[1:]:
                    valores = [
                        col.strip()
                        for col in linha_tabela.strip("|").split("|")
                    ]

                    # Ignora a linha separadora Markdown: |---|---|
                    if valores and all(
                        re.fullmatch(r":?-{3,}:?", valor.replace(" ", ""))
                        for valor in valores
                    ):
                        continue

                    linhas_dados.append(valores)

                tabela_doc = doc.add_table(
                    rows=1 + len(linhas_dados),
                    cols=len(cabecalho),
                )
                tabela_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
                tabela_doc.style = "Table Grid"
                aplicar_borda_tabela(tabela_doc)

                for col_idx, nome_col in enumerate(cabecalho):
                    cell = tabela_doc.cell(0, col_idx)
                    cell.text = nome_col

                    p_cell = cell.paragraphs[0]
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    for run in p_cell.runs:
                        run.bold = True
                        run.font.size = Pt(10)

                for row_idx, linha_dado in enumerate(linhas_dados):
                    for col_idx in range(len(cabecalho)):
                        valor = (
                            linha_dado[col_idx]
                            if col_idx < len(linha_dado)
                            else ""
                        )

                        cell = tabela_doc.cell(row_idx + 1, col_idx)
                        cell.text = valor

                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(9)

            continue

        elif linha.startswith("- ") or linha.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)

            texto_limpo = re.sub(r"^[-*]\s+", "", linha)
            partes = re.split(r"(\*\*.*?\*\*)", texto_limpo)

            for parte in partes:
                if parte.startswith("**") and parte.endswith("**"):
                    p.add_run(parte[2:-2]).bold = True
                else:
                    p.add_run(parte)

        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.paragraph_format.space_after = Pt(6)

            partes = re.split(r"(\*\*.*?\*\*)", linha)

            for parte in partes:
                if parte.startswith("**") and parte.endswith("**"):
                    p.add_run(parte[2:-2]).bold = True
                else:
                    p.add_run(parte)

        i += 1

    # Anexo com as figuras
    if figuras_geradas:
        p_sec_fig = doc.add_paragraph()
        p_sec_fig.paragraph_format.space_before = Pt(24)

        run_sec_fig = p_sec_fig.add_run(
            "ANEXO: ILUSTRAÇÕES E EVIDÊNCIAS GRÁFICAS"
        )
        run_sec_fig.bold = True
        run_sec_fig.font.size = Pt(12)

        for fig in figuras_geradas:
            caminho_img = fig.get("caminho")
            legenda = sanitizar_texto_latex(
                fig.get("legenda", "Ilustração técnica")
            )

            if caminho_img and Path(caminho_img).exists():
                p_leg = doc.add_paragraph()
                p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_leg.paragraph_format.space_before = Pt(14)
                p_leg.paragraph_format.space_after = Pt(4)

                run_l = p_leg.add_run(legenda)
                run_l.bold = True
                run_l.font.size = Pt(10)

                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(
                    str(caminho_img),
                    width=Cm(15),
                )

                p_fonte = doc.add_paragraph()
                p_fonte.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_fonte.paragraph_format.space_after = Pt(12)

                run_f = p_fonte.add_run(
                    "Fonte: Elaborado com base nos dados meteorológicos utilizados."
                )
                run_f.font.size = Pt(10)
                run_f.italic = True

    doc.save(str(caminho_docx_saida))


# ============================================================
# ORQUESTRADOR PRINCIPAL
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "Agente Editorial e Redator de Relatório Estatístico Climatológico"
        )
    )
    parser.parse_args()

    inicio_execucao = datetime.now()
    timestamp_str = inicio_execucao.strftime("%Y%m%d_%H%M%S")

    print("=" * 75)
    print("ETAPA 3: AGENTE EDITORIAL E REDATOR CIENTÍFICO")
    print("=" * 75)

    for arquivo in (
        ARQUIVO_DADOS_CSV,
        ARQUIVO_ESTATISTICAS,
        ARQUIVO_PROMPT,
    ):
        if not arquivo.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")

    # 1. Carregamento dos dados
    df = pd.read_csv(ARQUIVO_DADOS_CSV)

    if df.empty:
        raise ValueError("O CSV de coleta está vazio.")

    estatisticas_json = json.loads(
        ARQUIVO_ESTATISTICAS.read_text(encoding="utf-8")
    )

    metadados = estatisticas_json.get("metadados", {})
    periodo = metadados.get("periodo", {})

    print("\nDados estatísticos carregados:")
    print(
        f"   Localidade: {metadados.get('cidade', 'Não informada')} | "
        f"Período: {periodo.get('inicio', '')} a {periodo.get('fim', '')}"
    )
    print(f"   Observações no CSV: {len(df)}")

    # 2. Preparação do prompt editorial
    print(f"\n[1/4] Consultando OpenCode: {MODELO_PADRAO}")

    template_prompt = ARQUIVO_PROMPT.read_text(encoding="utf-8")

    substituicoes = {
        "{cidade}": str(metadados.get("cidade", "Não informada")),
        "{periodo_inicio}": str(periodo.get("inicio", "")),
        "{periodo_fim}": str(periodo.get("fim", "")),
        "{tamanho_amostral}": str(
            metadados.get("tamanho_amostral", len(df))
        ),
        "{estatisticas_json}": json.dumps(
            estatisticas_json,
            indent=2,
            ensure_ascii=False,
        ),
    }

    prompt_renderizado = template_prompt

    for marcador, valor in substituicoes.items():
        prompt_renderizado = prompt_renderizado.replace(marcador, valor)

    resposta_modelo = consultar_modelo(prompt_renderizado)

    print("\n" + "=" * 70)
    print("DEBUG: RESPOSTA BRUTA DO OPENCODE")
    print("=" * 70)
    print(resposta_modelo[:6000])
    print("\n[FIM DA RESPOSTA BRUTA]")

    with open(
        "debug_resposta_modelo.txt",
        "w",
        encoding="utf-8"
    ) as arquivo:
        arquivo.write(resposta_modelo)

    # 3. Extração da resposta
    plano_editorial, texto_markdown = extrair_resposta_dupla(
        resposta_modelo
    )

    titulo_relatorio = plano_editorial.get(
        "titulo_relatorio",
        "Relatório Técnico-Científico de Estatística Aplicada ao Clima",
    )

    figuras_solicitadas = plano_editorial.get(
        "figuras_encomendadas",
        [],
    )

    print("\n[2/4] Plano editorial recebido:")
    print(f"      Título: {titulo_relatorio}")
    print(f"      Figuras solicitadas: {len(figuras_solicitadas)}")

    # 4. Geração das figuras
    figuras_geradas = []

    for idx, config_fig in enumerate(figuras_solicitadas, start=1):
        if not isinstance(config_fig, dict):
            print(f"      Aviso: configuração {idx} ignorada.")
            continue

        tipo = config_fig.get("tipo_ferramenta", "desconhecida")
        legenda = config_fig.get("legenda_abnt", f"Figura {idx}")

        print(f"      [{idx}] {tipo}: {legenda}")

        caminho_fig = despachar_geracao_figura(
            df,
            config_fig,
            idx,
            PASTA_GRAFICOS,
        )

        if caminho_fig and Path(caminho_fig).exists():
            figuras_geradas.append(
                {
                    "id": f"figura_{idx}",
                    "tipo": tipo,
                    "legenda": legenda,
                    "caminho": str(caminho_fig),
                }
            )

    # 5. Persistência dos arquivos
    print("\n[3/4] Salvando plano editorial e Markdown...")

    texto_markdown_limpo = sanitizar_texto_latex(texto_markdown)

    ARQUIVO_SAIDA_PLANO.write_text(
        json.dumps(
            plano_editorial,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    ARQUIVO_SAIDA_MD.write_text(
        texto_markdown_limpo,
        encoding="utf-8",
    )

    # 6. Exportação para HTML e Word
    print("\n[4/4] Gerando HTML e Word...")

    gerar_relatorio_html(
        texto_markdown_limpo,
        figuras_geradas,
        ARQUIVO_SAIDA_HTML,
        titulo_relatorio,
    )

    gerar_relatorio_docx(
        texto_markdown_limpo,
        figuras_geradas,
        ARQUIVO_SAIDA_DOCX,
        metadados,
        titulo_relatorio,
    )

    # 7. Registro de auditoria
    registro_run = {
        "timestamp": timestamp_str,
        "etapa": "03_relatorio",
        "modelo": MODELO_PADRAO,
        "metadados": metadados,
        "plano_editorial_ia": plano_editorial,
        "artefatos_produzidos": {
            "plano_editorial": str(ARQUIVO_SAIDA_PLANO.name),
            "relatorio_markdown": str(ARQUIVO_SAIDA_MD.name),
            "relatorio_html": str(ARQUIVO_SAIDA_HTML.name),
            "relatorio_docx": str(ARQUIVO_SAIDA_DOCX.name),
            "total_figuras_solicitadas": len(figuras_solicitadas),
            "total_figuras_geradas": len(figuras_geradas),
        },
    }

    arquivo_run = PASTA_RUNS / f"run_{timestamp_str}_relatorio.json"

    arquivo_run.write_text(
        json.dumps(registro_run, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("\n" + "=" * 75)
    print("ETAPA 3 CONCLUÍDA!")
    print("=" * 75)
    print(f"1. Markdown: {ARQUIVO_SAIDA_MD}")
    print(f"2. HTML:     {ARQUIVO_SAIDA_HTML}")
    print(f"3. Word:     {ARQUIVO_SAIDA_DOCX}")
    print(f"4. Figuras:  {len(figuras_geradas)}")
    print(f"5. Auditoria: {arquivo_run}")
    print("=" * 75)


if __name__ == "__main__":
    main()
